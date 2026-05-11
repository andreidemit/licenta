"""
Build a Word draft of the thesis from licenta.md.

The goal is not to replace final manual proofreading, but to give the thesis a
proper academic DOCX structure: A4 page, Word heading styles, tables, lists,
code blocks, header/footer, and a TOC field that can be updated in Word.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "licenta.md"
OUTPUT = ROOT / "licenta_draft.docx"


def set_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after in [
        ("Title", 18, 0, 12),
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 14, 14, 6),
        ("Heading 3", 12, 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        style = styles.add_style("Code Block", 1)
        style.font.name = "Courier New"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        style.font.size = Pt(9)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_after = Pt(3)


def add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Actualizați câmpul în Word: Ctrl+A, F9"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr_text, fld_sep, placeholder, fld_end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Pagina ")
    add_field(paragraph, "PAGE")


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = "Simulator Grid-Based pentru Navigare Sigură"
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header_p.runs:
        set_font(run, size=9, color=(90, 90, 90))

    footer_p = section.footer.paragraphs[0]
    add_page_number(footer_p)
    for run in footer_p.runs:
        set_font(run, size=9, color=(90, 90, 90))


def clean_inline(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("__", "")
    text = text.replace("`", "")
    return text.strip()


def add_formatted_paragraph(doc: Document, text: str, style: str | None = None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_font(run, name="Courier New", size=10)
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            set_font(run, italic=True)
        else:
            run = p.add_run(part)
            set_font(run)
    return p


def add_cover(doc: Document) -> None:
    lines = [
        ("Universitatea din București", 14, True),
        ("Facultatea de Matematică și Informatică", 14, True),
        ("Departamentul de Informatică", 13, False),
    ]
    for text, size, bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, size=size, bold=bold)

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Simulator Grid-Based pentru Evaluarea Strategiilor")
    p.add_run().add_break()
    p.add_run("de Navigare Sigură în Medii Generate Procedural")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Lucrare de Licență")
    set_font(run, size=14, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()

    for text in [
        "Autor: Andrei Demit",
        "Coordonator științific: Lect. univ. dr. Florentina Suter",
        "Specializarea: Informatică",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, size=12)

    for _ in range(7):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("București, 2026")
    set_font(run, size=12, bold=True)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_paragraph("Cuprins", style="Heading 1")
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u')
    doc.add_page_break()


def is_table_start(lines, index) -> bool:
    if index + 1 >= len(lines):
        return False
    return lines[index].strip().startswith("|") and re.match(r"^\s*\|?[\s:\-\|]+\|?\s*$", lines[index + 1])


def parse_table(lines, index):
    rows = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [clean_inline(cell) for cell in raw.split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(set(cell.replace(":", "").strip()) <= {"-"} for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = text
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_font(run, size=10, bold=(r_idx == 0))
    doc.add_paragraph()


def should_skip_manual_toc(lines, index) -> tuple[bool, int]:
    if lines[index].strip().lower() != "## cuprins":
        return False, index
    index += 1
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("## Capitolul") or stripped.startswith("# Capitolul"):
            break
        index += 1
    return True, index


def convert_markdown_lines(doc: Document, lines: list[str]) -> None:
    in_code = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        skip, new_i = should_skip_manual_toc(lines, i)
        if skip:
            i = new_i
            continue

        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Code Block")
            run = p.add_run(line)
            set_font(run, name="Courier New", size=9)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if is_table_start(lines, i):
            table_rows, i = parse_table(lines, i)
            add_table(doc, table_rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            text = clean_inline(heading.group(2))
            if text.startswith("Capitolul "):
                doc.add_page_break()
                level = 1
            doc.add_paragraph(text, style=f"Heading {level}")
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            run = p.add_run(clean_inline(bullet.group(1)))
            set_font(run)
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            run = p.add_run(clean_inline(numbered.group(1)))
            set_font(run)
            i += 1
            continue

        add_formatted_paragraph(doc, stripped)
        i += 1


def split_markdown(markdown: str) -> tuple[list[str], list[str]]:
    """Return front matter (rezumat/abstract) and chapter body, skipping manual TOC."""
    lines = markdown.splitlines()
    rezumat_i = next((i for i, line in enumerate(lines) if line.strip() == "## Rezumat"), 0)
    toc_i = next((i for i, line in enumerate(lines) if line.strip().lower() == "## cuprins"), None)
    chapter_i = next((i for i, line in enumerate(lines) if line.strip().startswith("## Capitolul")), None)

    if toc_i is None or chapter_i is None:
        return lines[rezumat_i:], []

    front = lines[rezumat_i:toc_i]
    body = lines[chapter_i:]
    return front, body


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    front, body = split_markdown(markdown)

    final_doc = Document()
    configure_document(final_doc)
    add_header_footer(final_doc)
    add_cover(final_doc)
    convert_markdown_lines(final_doc, front)
    add_toc(final_doc)
    convert_markdown_lines(final_doc, body)
    final_doc.core_properties.title = "Simulator Grid-Based pentru Evaluarea Strategiilor de Navigare Sigură"
    final_doc.core_properties.author = "Andrei Demit"
    final_doc.core_properties.subject = "Lucrare de licență"
    final_doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
